#!/usr/bin/env python3
"""Generate the fictional Starlan document set for the RAG demo.

All content is fictional. DOCX is created with python-docx. PDF entries are
converted offline from the DOCX source with the bundled LibreOffice Kit.

Usage:
    python scripts/generate_mock_docs.py
    python scripts/generate_mock_docs.py --only A01,B05
"""

from __future__ import annotations

import argparse
import hashlib
import json
import os
import shutil
import subprocess
from pathlib import Path
from typing import Any

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

REPO_ROOT = Path(__file__).resolve().parents[1]
DATA_DIR = REPO_ROOT / "data"
DEFAULT_WORK_DIR = Path(r"D:\RAG库\build\mock-docs")
NODE = Path(r"D:\deepseek harness\resources\runtime\primary-runtime\dependencies\node\bin\node.exe")
LIBREOFFICE_CLI = Path(
    r"D:\deepseek harness\resources\app.asar.unpacked\dsh\node_modules\@deepseek-ai\libreoffice-kit\lib\cli.js"
)

COMPANY = "星澜精密仪器（虚构）有限公司"
COMPANY_EN = "Starlan Precision (Fictional) Co., Ltd."
DISCLAIMER = (
    f"{COMPANY} · 本文档全部内容为虚构演示资料，禁止用于真实生产、检验、采购或安全决策。"
)
FONT = "Microsoft YaHei"

# ---------------------------------------------------------------------------
# python-docx helpers
# ---------------------------------------------------------------------------


def set_run_font(run, size: float | None = None, bold: bool = False, color: str | None = None) -> None:
    run.font.name = FONT
    run._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    if size:
        run.font.size = Pt(size)
    run.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_text(cell, text: str, bold: bool = False, size: float = 9.0, color: str | None = None) -> None:
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_after = Pt(0)
    run = paragraph.add_run(str(text))
    set_run_font(run, size=size, bold=bold, color=color)


def set_col_widths(table, widths_cm: list[float]) -> None:
    table.autofit = False
    for row in table.rows:
        for idx, width in enumerate(widths_cm):
            row.cells[idx].width = Cm(width)


def add_table(
    doc: Document,
    headers: list[str],
    rows: list[list[Any]],
    widths_cm: list[float] | None = None,
    caption: str | None = None,
) -> None:
    if caption:
        p = doc.add_paragraph()
        run = p.add_run(caption)
        set_run_font(run, size=9, bold=True, color="555555")
        p.paragraph_format.space_after = Pt(3)

    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    header = table.rows[0].cells
    for idx, text in enumerate(headers):
        set_cell_shading(header[idx], "D9EAF7")
        set_cell_text(header[idx], text, bold=True, size=9)
    for row in rows:
        cells = table.add_row().cells
        for idx, value in enumerate(row):
            set_cell_text(cells[idx], value, size=9)
    if widths_cm and len(widths_cm) == len(headers):
        set_col_widths(table, widths_cm)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def add_page_field(paragraph) -> None:
    run = paragraph.add_run()
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = "PAGE"
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._r.append(fld_begin)
    run._r.append(instr)
    run._r.append(fld_end)
    set_run_font(run, size=8, color="777777")


def add_header_footer(doc: Document, spec: dict[str, Any]) -> None:
    section = doc.sections[0]
    header = section.header
    hp = header.paragraphs[0]
    hp.text = ""
    hp.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = hp.add_run(f"{COMPANY}    {spec['code']}    {spec['version']}")
    set_run_font(run, size=8, color="666666")
    p_pr = hp._p.get_or_add_pPr()
    p_bdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "4")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "BFBFBF")
    p_bdr.append(bottom)
    p_pr.append(p_bdr)

    footer = section.footer
    fp = footer.paragraphs[0]
    fp.text = ""
    fp.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = fp.add_run(DISCLAIMER)
    set_run_font(run, size=7.5, color="777777")

    page_p = footer.add_paragraph()
    page_p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = page_p.add_run("第 ")
    set_run_font(run, size=8, color="777777")
    add_page_field(page_p)
    run = page_p.add_run(" 页")
    set_run_font(run, size=8, color="777777")


def add_control_table(doc: Document, spec: dict[str, Any]) -> None:
    rows = [
        ["文件编号", spec["code"], "版本", spec["version"]],
        ["发布日期", spec["date"], "文档状态", spec["status"]],
        ["适用产品/范围", spec["scope"], "密级", "内部演示（虚构）"],
    ]
    table = doc.add_table(rows=0, cols=4)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for row in rows:
        cells = table.add_row().cells
        set_cell_shading(cells[0], "F2F2F2")
        set_cell_shading(cells[2], "F2F2F2")
        set_cell_text(cells[0], row[0], bold=True, size=9)
        set_cell_text(cells[1], row[1], size=9)
        set_cell_text(cells[2], row[2], bold=True, size=9)
        set_cell_text(cells[3], row[3], size=9)
    set_col_widths(table, [3.2, 6.4, 3.2, 4.0])
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def add_blocks(doc: Document, blocks: list[tuple[Any, ...]]) -> None:
    for block in blocks:
        kind = block[0]
        if kind == "h1":
            p = doc.add_heading(block[1], level=1)
            for run in p.runs:
                set_run_font(run, size=14, bold=True, color="1F3864")
        elif kind == "h2":
            p = doc.add_heading(block[1], level=2)
            for run in p.runs:
                set_run_font(run, size=11.5, bold=True, color="2E5C8A")
        elif kind == "p":
            p = doc.add_paragraph()
            run = p.add_run(block[1])
            set_run_font(run, size=10)
            p.paragraph_format.first_line_indent = Cm(0.74)
            p.paragraph_format.line_spacing = 1.35
        elif kind == "bullets":
            for item in block[1]:
                p = doc.add_paragraph(style="List Bullet")
                run = p.add_run(item)
                set_run_font(run, size=10)
                p.paragraph_format.line_spacing = 1.25
        elif kind == "note":
            p = doc.add_paragraph()
            run = p.add_run(block[1])
            set_run_font(run, size=9, color="8B0000" if block[2:] == ("red",) else "555555")
            p.paragraph_format.left_indent = Cm(0.5)
        elif kind == "table":
            _, headers, rows, widths = block[:4]
            caption = block[4] if len(block) > 4 else None
            add_table(doc, headers, rows, widths, caption)
        elif kind == "pagebreak":
            doc.add_page_break()
        else:
            raise ValueError(f"unknown block kind: {kind}")


def new_document(spec: dict[str, Any]) -> Document:
    doc = Document()
    section = doc.sections[0]
    section.page_height = Cm(29.7)
    section.page_width = Cm(21.0)
    section.top_margin = Cm(2.2)
    section.bottom_margin = Cm(2.0)
    section.left_margin = Cm(2.2)
    section.right_margin = Cm(2.0)

    normal = doc.styles["Normal"]
    normal.font.name = FONT
    normal.font.size = Pt(10)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    for level in (1, 2, 3):
        style = doc.styles[f"Heading {level}"]
        style.font.name = FONT
        style._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)

    add_header_footer(doc, spec)
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run(spec["title"])
    set_run_font(run, size=20, bold=True, color="1F3864")
    title.paragraph_format.space_after = Pt(4)

    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = sub.add_run(f"{COMPANY_EN}  ·  {spec['category']}")
    set_run_font(run, size=9, color="666666")
    sub.paragraph_format.space_after = Pt(10)

    if spec.get("obsolete"):
        warn = doc.add_paragraph()
        warn.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = warn.add_run("【已废止，仅供存档】")
        set_run_font(run, size=17, bold=True, color="C00000")
        warn.paragraph_format.space_after = Pt(8)

    add_control_table(doc, spec)
    for block in spec.get("blocks", []):
        if block[0] == "pagebreak":
            doc.add_page_break()
        else:
            add_blocks(doc, [block])
    doc.core_properties.title = spec["title"]
    doc.core_properties.author = COMPANY
    doc.core_properties.subject = "虚构演示资料"
    doc.core_properties.comments = "全部内容为虚构，禁止用于真实业务。"
    return doc


# ---------------------------------------------------------------------------
# Document content
# ---------------------------------------------------------------------------

X100_COMMON = [
    ("h1", "1. 产品概述"),
    ("p", "X-100 便携式扭矩校验仪用于现场校验扭矩扳手、电动拧紧工具与扭矩传感器。设备采用铝合金一体化外壳，内置 3.7 V 锂离子电池，适合产线巡检与实验室抽检。"),
    ("p", "本机通过 2.4 英寸彩色液晶屏显示实时扭矩、峰值扭矩与单位，支持峰值保持、实时跟踪和五点校准。随机提供铝合金携带箱、充电器与 USB-C 数据线。"),
    ("h1", "2. 技术参数"),
    ("table", ["参数", "X-100 规格"], [
        ["测量范围", "0.5–100 N·m"],
        ["示值精度", "±0.5% FS"],
        ["分辨率", "0.01 N·m"],
        ["单位", "N·m / kgf·cm / lbf·ft"],
        ["额定功率", "5 W"],
        ["防护等级", "IP54"],
        ["通信方式", "Bluetooth 5.0"],
        ["电池", "3.7 V 2600 mAh 锂离子电池"],
        ["工作温度", "-10 °C ~ 50 °C"],
        ["外形尺寸", "210 × 95 × 42 mm"],
        ["重量", "约 0.68 kg（不含携带箱）"],
    ], [4.8, 8.5], "表 1  X-100 主要技术参数"),
    ("h1", "3. 结构、按键与接口"),
    ("p", "整机由显示单元、传感器接口、按键区、USB-C 充电口和腕带挂点组成。按键区自左至右为电源键、单位键、模式键、清零键和确认键。长按电源键 2 s 开机或关机。"),
    ("bullets", [
        "传感器接口：标准 1/4 英寸方头，可更换 3/8 英寸转接头。",
        "USB-C：仅用于充电与本地数据导出，不支持边充边用。",
        "腕带挂点：用于防跌落腕带，现场使用时应佩戴。",
    ]),
    ("h1", "4. 操作步骤"),
    ("p", "开机后先选择单位与模式，再安装对应转接头。空载状态下按清零键，屏幕显示 0.00 后开始加载。测量峰值时选择 PEAK 模式；需要连续观察时选择 TRACK 模式。"),
    ("bullets", [
        "步骤 1：检查外壳、屏幕与传感器接口无损伤。",
        "步骤 2：长按电源键开机，确认电量高于 20%。",
        "步骤 3：选择 N·m 单位与 PEAK 模式，空载清零。",
        "步骤 4：缓慢加载至目标扭矩，读取峰值并记录。",
        "步骤 5：卸载后按清零键，关闭电源并放入携带箱。",
    ]),
    ("h1", "5. 蓝牙连接与数据导出"),
    ("p", "X-100 通过 Bluetooth 5.0 与手机 App 或桌面工具连接。开机后蓝牙处于可发现状态，连接成功后屏幕右上角显示蓝牙图标。单次最多传输 500 条记录，超出时按时间顺序覆盖最旧记录。"),
    ("p", "导出的 CSV 文件包含序号、时间、单位、峰值、模式和设备序列号。设备序列号可在电池仓标签上查看，导出前应核对。"),
    ("h1", "6. 校准与维护"),
    ("p", "建议校准周期为 12 个月，或累计使用 5000 次后提前校准。校准应在 20 ± 5 °C 环境内，使用不确定度优于 0.1% 的标准扭矩源，至少取量程的 20%、40%、60%、80%、100% 五个点。"),
    ("p", "外壳可用微湿无绒布擦拭，禁止使用丙酮、香蕉水或含研磨颗粒的清洁剂。传感器接口应保持干燥，不得加注润滑油。"),
    ("h1", "7. 常见故障与处理"),
    ("table", ["现象", "可能原因", "处理"], [
        ["屏幕不亮", "电量耗尽或电源键未长按", "充电 30 min 后长按电源键 2 s"],
        ["数值漂移", "未空载清零或接口松动", "重新安装转接头并空载清零"],
        ["蓝牙无法连接", "距离过远或 App 权限未开启", "保持在 10 m 内，检查蓝牙权限"],
        ["充电无反应", "充电器或线缆异常", "更换 5 V/1 A 充电器与线缆"],
    ], [4.0, 5.5, 5.5], "表 2  X-100 常见故障"),
    ("h1", "8. 安全声明"),
    ("p", "本设备仅用于扭矩测量，不得作为起重、牵引或承力结构使用。加载时操作人员应佩戴护目镜，避免站在工具回弹方向。电池鼓包、漏液或外壳破损时应停止使用。"),
]

X100S_BLOCKS = [
    ("h1", "1. 产品概述"),
    ("p", "X-100S 便携式扭矩校验仪面向需要无线数据回传与更高防护等级的现场场景。整机沿用 X-100 的铝合金机身、2.4 英寸显示屏与 1/4 英寸传感器接口，外形尺寸和按键布局保持一致。"),
    ("p", "X-100S 在便携平台上增加了 Wi-Fi 6 数据通道与更高等级的密封结构，适合油污、粉尘和冲洗环境下的巡检作业。设备同时保留 Bluetooth 5.2，用于手机 App 近场连接。"),
    ("h1", "2. 技术参数"),
    ("table", ["参数", "X-100S 规格"], [
        ["测量范围", "0.5–100 N·m"],
        ["示值精度", "±0.5% FS"],
        ["分辨率", "0.01 N·m"],
        ["单位", "N·m / kgf·cm / lbf·ft"],
        ["额定功率", "8 W"],
        ["防护等级", "IP67"],
        ["通信方式", "Wi-Fi 6（802.11ax）+ Bluetooth 5.2，支持 MQTT"],
        ["电池", "3.7 V 2600 mAh 锂离子电池"],
        ["工作温度", "-10 °C ~ 50 °C"],
        ["外形尺寸", "210 × 95 × 42 mm"],
        ["重量", "约 0.68 kg（不含携带箱）"],
    ], [4.8, 8.5], "表 1  X-100S 主要技术参数"),
    ("h1", "3. 结构与接口"),
    ("p", "X-100S 的显示单元、传感器接口、按键区、USB-C 充电口和腕带挂点与 X-100 相同。机身右侧增加 Wi-Fi 状态指示灯，蓝色常亮表示已接入企业无线网络，蓝色闪烁表示正在配网。"),
    ("p", "首次配网时，长按模式键 5 s 进入配网模式，在 App 中选择“X-100S 设备网络”并输入企业 Wi-Fi 凭据。设备仅支持 2.4 GHz/5 GHz 双频网络，不支持强制门户认证。"),
    ("h1", "4. 操作步骤"),
    ("p", "开机后选择单位与模式，空载清零。测量峰值时选择 PEAK 模式；需要连续观察时选择 TRACK 模式。Wi-Fi 连接成功后，测量记录可按批次上传至数据平台。"),
    ("bullets", [
        "步骤 1：检查密封圈、屏幕与传感器接口无损伤。",
        "步骤 2：长按电源键开机，确认电量高于 20%。",
        "步骤 3：选择 N·m 单位与 PEAK 模式，空载清零。",
        "步骤 4：缓慢加载至目标扭矩，读取峰值并记录。",
        "步骤 5：在 App 中选择批次并上传；确认上传成功后再关机。",
    ]),
    ("h1", "5. 数据通信"),
    ("p", "X-100S 支持 MQTT 主题上报，主题格式为 starlan/{site}/{deviceId}/torque。默认上报间隔为 5 s，断网时本地缓存最近 2000 条记录，恢复网络后按时间顺序补传。"),
    ("p", "Wi-Fi 与蓝牙可同时开启，但同一时刻只允许一个 App 建立控制连接。数据平台下发的时间同步命令会在每次连接建立后执行一次。"),
    ("h1", "6. 校准与维护"),
    ("p", "建议校准周期为 12 个月，或累计使用 5000 次后提前校准。校准方法、环境条件与标准扭矩源要求与 X-100 相同。每次校准后应检查密封圈是否老化，必要时更换。"),
    ("p", "外壳可用清水冲洗，但冲洗前必须确认 USB-C 防水塞完全闭合。禁止使用高压水枪直射传感器接口和显示屏边缘。"),
    ("h1", "7. 常见故障与处理"),
    ("table", ["现象", "可能原因", "处理"], [
        ["Wi-Fi 无法配网", "网络为强制门户或 5 GHz 信道不支持", "改用 2.4 GHz 网络或联系 IT"],
        ["上传失败", "MQTT 凭据过期或网络中断", "检查凭据；恢复网络后等待补传"],
        ["密封报警", "防水塞未闭合或密封圈老化", "闭合防水塞；更换密封圈"],
        ["蓝牙无法连接", "已被 Wi-Fi 控制连接占用", "断开 Wi-Fi 控制连接后重试"],
    ], [4.0, 5.5, 5.5], "表 2  X-100S 常见故障"),
    ("h1", "8. 安全声明"),
    ("p", "本设备仅用于扭矩测量，不得作为起重、牵引或承力结构使用。充电时必须打开 USB-C 防水塞并在干燥环境进行；充电完成后立即闭合防水塞。"),
]

X200_BLOCKS = [
    ("h1", "1. 产品概述"),
    ("p", "X-200 台式多通道扭矩校验系统面向实验室、计量室和产线固定工位。系统采用与 X 系列一致的深灰机身、蓝色装饰条和前置接口布局，但定位为固定式多通道校验平台，不提供电池供电或便携携带箱。"),
    ("p", "X-200 配备 12 英寸触摸屏、四路独立传感器输入和工业以太网接口，可同时校验四把扭矩工具。系统内置工装夹具安装孔，建议固定在承重不小于 80 kg 的实验台上。"),
    ("h1", "2. 技术参数"),
    ("table", ["参数", "X-200 规格"], [
        ["测量范围", "5–2000 N·m（四通道独立）"],
        ["示值精度", "±0.2% FS"],
        ["分辨率", "0.1 N·m"],
        ["通道数", "4"],
        ["额定功率", "300 W"],
        ["防护等级", "IP20"],
        ["通信方式", "Ethernet / RS-485 / Modbus TCP"],
        ["供电", "220 V AC，50 Hz"],
        ["显示", "12 英寸工业触摸屏"],
        ["外形尺寸", "420 × 360 × 180 mm"],
        ["重量", "约 8.5 kg"],
    ], [4.8, 8.5], "表 1  X-200 主要技术参数"),
    ("h1", "3. 安装与接口"),
    ("p", "将系统放置在水平实验台上，使用 M8 螺栓固定四角安装孔。电源接入前确认接地电阻小于 4 Ω。四个传感器接口位于机身背面，标识为 CH1 至 CH4；以太网口位于右侧，默认地址为 192.168.10.20。"),
    ("p", "首次上电后系统执行自检，检查通道零点、触摸屏、风扇和通信模块。自检通过后进入主界面；若任一通道零点漂移超过 0.05% FS，应重新预热 30 min 后再试。"),
    ("h1", "4. 操作步骤"),
    ("p", "在主界面选择通道与单位，安装对应工装，空载清零。选择“单次校验”或“连续校验”模式，设置目标扭矩、采样率和判定上下限。校验完成后可导出 CSV 或 PDF 报告。"),
    ("bullets", [
        "步骤 1：确认实验台固定、接地可靠、通道为空载。",
        "步骤 2：上电预热 30 min，执行通道自检。",
        "步骤 3：安装工装与传感器，选择通道并清零。",
        "步骤 4：设置目标值与判定限，开始校验。",
        "步骤 5：保存记录，导出报告，关闭系统电源。",
    ]),
    ("h1", "5. 报告与数据管理"),
    ("p", "系统按“日期/班次/操作员”建立目录，单次校验报告包含通道号、目标扭矩、峰值、示值误差、重复性和判定结果。历史数据默认保留 24 个月。"),
    ("p", "导出 PDF 报告时使用系统内置模板，报告页脚自动带出设备序列号与校准有效期。不得手工修改报告中的原始测量值。"),
    ("h1", "6. 校准与维护"),
    ("p", "建议校准周期为 12 个月。校准时应使用不确定度优于 0.05% 的标准扭矩源，对每个通道分别取 20%、40%、60%、80%、100% 五个点。长期停用后重新启用应先预热 60 min。"),
    ("p", "每季度清理风扇滤网，每年检查接地线与电源线。禁止在带电状态下插拔传感器连接器。"),
    ("h1", "7. 安全声明"),
    ("p", "X-200 为固定式设备，必须由培训合格人员操作。加载时操作人员不得站在工装回弹方向；急停按钮位于操作台右下角，按下后切断伺服与加热输出。"),
]

QS_COMMON_HEADER = [
    ("h1", "1. 目的与范围"),
    ("p", "本规范规定 X-100 与 X-100S 便携式扭矩校验仪成品出厂检验的项目、条件、抽样与判定规则。适用于星澜精密仪器（虚构）有限公司质量部及授权检验人员。"),
    ("h1", "2. 引用文件"),
    ("p", "下列文件中的条款通过本规范引用而成为本规范的条款。凡是注日期的引用文件，仅注日期对应的版本适用；凡是不注日期的引用文件，其最新版本适用。"),
    ("bullets", [
        "《环境试验细则》（资料库未收录，本文仅作为引用文件列出）",
        "《成品检验抽样规定》Q/STARLAN-QA-002",
        "《扭矩计量器具校准规范》Q/STARLAN-MET-001",
    ]),
    ("h1", "3. 检验条件"),
    ("p", "检验环境温度应为 20 ± 5 °C，相对湿度 45%–75%，无强电磁干扰和明显振动。检验前设备应在环境中静置不少于 2 h，标准扭矩源应处于有效校准期内。"),
]

QS20_ROWS = [
    ["1", "外观", "表面无裂纹、锈蚀、明显划痕；标识清晰", "目视", "全检"],
    ["2", "外形尺寸", "210 × 95 × 42 mm，允许偏差 ±0.5 mm", "游标卡尺", "AQL 1.0"],
    ["3", "整机重量", "0.68 ± 0.03 kg（不含携带箱）", "电子秤", "AQL 1.0"],
    ["4", "扭矩示值误差", "≤ ±0.5% FS", "标准扭矩仪，五点", "每批 3 台"],
    ["5", "重复性", "≤ 0.3%", "同一扭矩点重复 10 次", "每批 3 台"],
    ["6", "绝缘电阻", "≥ 100 MΩ（500 VDC）", "绝缘电阻测试仪", "全检"],
    ["7", "耐压", "1500 VAC / 60 s，无击穿、无飞弧", "耐压测试仪", "全检"],
    ["8", "防护等级", "IP54", "按 GB/T 4208 方法", "每批 1 台"],
]

QS21_ROWS = [
    ["1", "外观", "表面无裂纹、锈蚀、明显划痕；标识清晰", "目视", "全检"],
    ["2", "外形尺寸", "210 × 95 × 42 mm，允许偏差 ±0.5 mm", "游标卡尺", "AQL 1.0"],
    ["3", "整机重量", "0.68 ± 0.03 kg（不含携带箱）", "电子秤", "AQL 1.0"],
    ["4", "扭矩示值误差", "≤ ±0.4% FS", "标准扭矩仪，五点", "每批 3 台"],
    ["5", "重复性", "≤ 0.2%", "同一扭矩点重复 10 次", "每批 3 台"],
    ["6", "绝缘电阻", "≥ 500 MΩ（500 VDC）", "绝缘电阻测试仪", "全检"],
    ["7", "耐压", "2000 VAC / 60 s，无击穿、无飞弧", "耐压测试仪", "全检"],
    ["8", "防护等级", "X-100：IP54；X-100S：IP67", "按 GB/T 4208 方法", "每批 1 台"],
]

QS20_BLOCKS = QS_COMMON_HEADER + [
    ("h1", "4. 检验项目与要求"),
    ("p", "下表为 V2.0 版本的检验项目。表中数值为本版本的有效要求，检验记录应注明依据版本。"),
    ("table", ["序号", "项目", "要求", "方法", "抽样/频次"], QS20_ROWS, [1.2, 2.8, 6.2, 3.6, 2.8], "表 1  成品检验项目（V2.0）"),
    ("h1", "5. 抽样与判定"),
    ("p", "外观、绝缘电阻和耐压为关键项，任一不合格则整批拒收。一般项按 AQL 1.0 抽样，Ac 值按《成品检验抽样规定》执行；一般项不合格数超过 Ac 值时整批拒收。"),
    ("p", "检验记录保存 5 年。记录内容包括批次号、序列号、检验日期、检验员、标准器编号、环境温湿度、实测值和判定结论。"),
    ("h1", "6. 版本与修订"),
    ("p", "V2.0 于 2025-03-10 发布并实施，替代 V1.4。V2.0 将耐压测试电压定为 1500 VAC，绝缘电阻下限定为 100 MΩ。后续版本如变更上述指标，应在本规范中明确标注版本和发布日期。"),
    ("h1", "7. 记录表单"),
    ("p", "检验记录使用《X-100 系列成品检验记录表》QR-QA-011。记录不得使用铅笔填写；更正时应划改并签名，不得涂改。"),
]

QS21_BLOCKS = QS_COMMON_HEADER + [
    ("h1", "4. 检验项目与要求"),
    ("p", "下表为 V2.1 版本的检验项目。V2.1 于 2026-08-15 发布，替代 V2.0；检验记录应注明依据版本，避免与历史批次混淆。"),
    ("table", ["序号", "项目", "要求", "方法", "抽样/频次"], QS21_ROWS, [1.2, 2.8, 6.2, 3.6, 2.8], "表 1  成品检验项目（V2.1）"),
    ("h1", "5. 抽样与判定"),
    ("p", "外观、绝缘电阻和耐压为关键项，任一不合格则整批拒收。一般项按 AQL 1.0 抽样。V2.1 将扭矩示值误差收紧至 ±0.4% FS，重复性收紧至 0.2%，并提高了绝缘电阻与耐压要求。"),
    ("p", "检验记录保存 5 年。记录应包含依据版本、批次号、序列号、标准器编号、环境温湿度和实测值。"),
    ("h1", "6. 版本与修订记录"),
    ("table", ["版本", "日期", "主要变更"], [
        ["V2.0", "2025-03-10", "耐压 1500 VAC；绝缘电阻 ≥100 MΩ；误差 ±0.5% FS；重复性 0.3%"],
        ["V2.1", "2026-08-15", "耐压 2000 VAC；绝缘电阻 ≥500 MΩ；误差 ±0.4% FS；重复性 0.2%；X-100S 防护等级 IP67"],
    ], [2.5, 3.5, 10.5], "表 2  版本修订记录"),
    ("h1", "7. 记录表单"),
    ("p", "V2.1 使用《X-100 系列成品检验记录表》QR-QA-011 的 2026 版。旧版记录表可用于历史批次追溯，不得用于新批次。"),
]

M3_BLOCKS = [
    ("h1", "1. 安全警告"),
    ("p", "M-3 半自动扭矩装配台用于金属壳体螺钉锁付。操作人员必须经过培训并佩戴护目镜；长发应束起，不得佩戴手套操作旋转部件。维修前必须执行挂牌上锁（LOTO）并等待主轴完全停止。"),
    ("h1", "2. 设备组成"),
    ("p", "设备由机架、伺服拧紧轴、三轴移动平台、工装夹具、7 英寸触摸屏、双手启动按钮、急停按钮和安全光栅组成。气源接口位于机架左侧，要求 0.5–0.7 MPa 干燥压缩空气。"),
    ("h1", "3. 技术参数"),
    ("table", ["参数", "M-3 规格"], [
        ["默认锁付扭矩", "8.0 N·m"],
        ["转速", "300 rpm"],
        ["单件节拍", "≤ 12 s"],
        ["重复精度", "±1%"],
        ["程序容量", "50 组"],
        ["电源", "220 V AC，50 Hz，1.5 kW"],
        ["通信", "RS-485"],
        ["外形尺寸", "1200 × 800 × 1600 mm"],
    ], [5.0, 8.0], "表 1  M-3 主要参数"),
    ("h1", "4. 开机与回原点"),
    ("bullets", [
        "步骤 1：检查安全光栅、急停按钮和双手按钮功能正常。",
        "步骤 2：打开气源，确认压力在 0.5–0.7 MPa。",
        "步骤 3：合上主电源，等待触摸屏进入主界面。",
        "步骤 4：点击“回原点”，确认三轴依次回到机械零点。",
        "步骤 5：调用程序，空运行一次确认轨迹无干涉。",
    ]),
    ("h1", "5. 程序与参数"),
    ("p", "在主界面选择程序号，设置目标扭矩、转速、缓启动时间和锁付点数。M-3 默认锁付扭矩为 8.0 N·m，转速 300 rpm，缓启动时间 0.3 s。修改参数后应保存并记录变更人。"),
    ("h1", "6. 换型步骤"),
    ("p", "从 NX-100 换型到 NX-200 时，先更换工装夹具与定位销，再调用 NX-200 程序，最后用首件确认锁付位置和扭矩。换型后必须执行首件检验，合格后方可批量生产。"),
    ("h1", "7. 报警与处理"),
    ("table", ["代码", "含义", "处理"], [
        ["E01", "急停按下", "排除危险后旋起急停，复位并重新回原点"],
        ["E03", "安全光栅遮挡", "清除遮挡物，按复位按钮"],
        ["E05", "扭矩超上限", "检查程序、批头与螺纹，重新首件确认"],
        ["E08", "气源压力低", "检查气源与过滤减压阀"],
    ], [2.0, 5.0, 8.0], "表 2  常见报警"),
    ("h1", "8. 日常保养"),
    ("p", "每班清洁台面与夹具，每周检查导轨润滑与气管，每月检查安全光栅与急停回路。扭矩校准周期为 6 个月；更换伺服轴或批头后应重新校准。"),
]

M3PRO_BLOCKS = [
    ("h1", "1. 安全警告"),
    ("p", "M-3 Pro 全自动扭矩装配台用于高节拍螺钉锁付。设备运行时安全门必须关闭；打开安全门会触发停机。维修前执行 LOTO，并等待主轴与三轴平台完全停止。"),
    ("h1", "2. 设备组成与定位"),
    ("p", "M-3 Pro 由伺服拧紧轴、四轴移动平台、双工位转台、10 英寸触摸屏、安全门锁、急停按钮和工业以太网接口组成。设备定位为全自动高节拍装配，不提供手动双手启动模式。"),
    ("p", "本机程序模板可兼容 X-900 扩展型号，该型号不在本手册范围内。"),
    ("h1", "3. 技术参数"),
    ("table", ["参数", "M-3 Pro 规格"], [
        ["默认锁付扭矩", "12.0 N·m"],
        ["转速", "450 rpm"],
        ["单件节拍", "≤ 8 s"],
        ["重复精度", "±0.5%"],
        ["程序容量", "200 组"],
        ["电源", "220 V AC，50 Hz，2.2 kW"],
        ["通信", "Ethernet / Modbus TCP"],
        ["外形尺寸", "1500 × 1000 × 1800 mm"],
    ], [5.0, 8.0], "表 1  M-3 Pro 主要参数"),
    ("h1", "4. 开机与自动运行"),
    ("bullets", [
        "步骤 1：确认安全门关闭、急停释放、气源压力正常。",
        "步骤 2：合上主电源，等待控制器自检完成。",
        "步骤 3：选择程序，执行回原点与空运行。",
        "步骤 4：放入首件，执行首件锁付并确认扭矩曲线。",
        "步骤 5：切换自动模式，启动连续生产。",
    ]),
    ("h1", "5. 程序与换型"),
    ("p", "M-3 Pro 默认锁付扭矩为 12.0 N·m，转速 450 rpm，缓启动 0.2 s。换型时更换快换夹具与批头，调用对应程序，再用首件确认锁付深度、扭矩和节拍。程序切换后应核对程序号与产品型号。"),
    ("h1", "6. 报警与处理"),
    ("table", ["代码", "含义", "处理"], [
        ["E01", "安全门未关闭", "关闭安全门并复位"],
        ["E05", "扭矩超上限", "检查程序、批头和螺纹；必要时更换批头"],
        ["E09", "转台定位超时", "检查转台传感器与异物"],
        ["E12", "以太网通信中断", "检查网线与上位机地址"],
    ], [2.0, 5.0, 8.0], "表 2  常见报警"),
    ("h1", "7. 维护与校准"),
    ("p", "每班清洁夹具与转台，每周检查安全门锁与光栅，每月检查导轨润滑和伺服电缆。扭矩校准周期为 6 个月；更换伺服轴、批头或夹具后应重新校准并做首件确认。"),
]

MAINT_BLOCKS = [
    ("h1", "1. 目的与范围"),
    ("p", "本规程规定 M-3、M-3 Pro 扭矩装配台的日常点检、周期保养、扭矩校准与记录要求，适用于设备部、生产部与质量部。"),
    ("h1", "2. 职责"),
    ("bullets", [
        "操作员：每班执行班前点检与清洁，发现异常立即停机并报告。",
        "设备技术员：执行周/月保养，处理一般故障。",
        "计量室：执行扭矩校准并出具证书。",
        "质量部：监督保养记录与校准有效期。",
    ]),
    ("h1", "3. 保养周期与项目"),
    ("table", ["周期", "项目", "标准", "记录"], [
        ["每班", "台面、夹具、批头清洁", "无油污、无金属屑", "班前点检表"],
        ["每班", "安全光栅、急停、双手按钮", "功能正常", "班前点检表"],
        ["每周", "导轨润滑、气管泄漏", "无干摩擦、无泄漏", "周保养记录"],
        ["每月", "安全门锁、伺服电缆、接地", "无松动、接地 <4 Ω", "月保养记录"],
        ["每 500 h", "减速机润滑脂", "按牌号补充", "润滑记录"],
        ["每 6 个月", "扭矩校准", "五点校准合格", "校准证书"],
    ], [2.5, 6.0, 5.0, 3.5], "表 1  保养周期"),
    ("h1", "4. 扭矩校准"),
    ("p", "校准使用不确定度优于 0.1% 的标准扭矩源，至少取 20%、40%、60%、80%、100% 五个点，每点重复 3 次。示值误差超过 ±1% 或重复性超过 0.5% 时，应调整或维修后重新校准。"),
    ("p", "校准周期为 6 个月，或更换伺服轴、批头、夹具后立即校准。校准记录应包含设备编号、标准器编号、环境温度、实测值、结论和有效期。"),
    ("h1", "5. 停机与恢复"),
    ("p", "长期停机前应清洁设备、释放气源、断开电源并覆盖防尘罩。恢复生产前执行空运行、安全回路测试和首件确认。"),
]

SAFE_ASSEMBLY_BLOCKS = [
    ("h1", "1. 适用范围"),
    ("p", "本规程适用于 M-3、M-3 Pro 装配工位的日常操作、换型、清洁和异常处理。所有进入工位的人员必须遵守本规程。"),
    ("h1", "2. 个人防护"),
    ("bullets", [
        "佩戴护目镜，防止批头或螺钉飞溅。",
        "长发束起，不得佩戴手套操作旋转部件。",
        "穿防砸安全鞋，不得穿拖鞋或凉鞋进入工位。",
        "更换批头时使用专用工具，不得徒手施力。",
    ]),
    ("h1", "3. 开机前检查"),
    ("p", "检查安全门、光栅、急停按钮、双手按钮和气源压力。确认工装夹具无裂纹、定位销无磨损、批头无崩刃。任何一项异常均不得开机。"),
    ("h1", "4. 紧急停机与恢复"),
    ("p", "发生人身危险、设备异响、冒烟或工件飞出时，立即按下急停按钮。急停后必须等待主轴完全停止（约 3 s），确认危险消除后方可旋起急停、复位安全回路并重新回原点。"),
    ("p", "急停恢复后，第一件产品必须执行首件确认；未经确认不得直接批量生产。"),
    ("h1", "5. 换型安全"),
    ("p", "换型前切断伺服使能并执行 LOTO。更换工装、夹具和批头后，手动盘车确认无干涉，再低速空运行。首件确认合格后方可切换自动模式。"),
    ("h1", "6. 异常与报告"),
    ("p", "发现安全光栅失效、急停按钮卡滞、安全门锁损坏时，立即停机并悬挂“禁止操作”标识，通知设备部处理。不得短接或旁路安全回路。"),
]

SAFE_ELECTRICAL_BLOCKS = [
    ("h1", "1. 适用范围"),
    ("p", "本规程适用于 X-100、X-100S 成品耐压、绝缘电阻和接地连续性测试。测试区域应设置黄色警戒线和“高压危险”标识。"),
    ("h1", "2. 测试前准备"),
    ("bullets", [
        "确认耐压测试仪、绝缘电阻测试仪在有效校准期内。",
        "检查测试线绝缘层无破损，接地夹连接可靠。",
        "测试人员穿绝缘鞋，佩戴护目镜，不得佩戴金属饰品。",
        "测试区域只允许一名操作员和一名监护人在场。",
    ]),
    ("h1", "3. 测试步骤"),
    ("p", "先将测试仪接地端连接到产品接地端子，再连接高压端。设置测试电压与时间，确认无人接触产品后启动测试。测试过程中监护人不得离开。"),
    ("p", "耐压测试结束后，必须先切断高压输出，并使用放电棒对被测端放电不少于 5 分钟，确认无残余电压后方可拆线。"),
    ("h1", "4. 判定与记录"),
    ("p", "耐压测试无击穿、无飞弧为合格；绝缘电阻不低于当前版本规范要求为合格。测试记录应包含产品序列号、测试电压、时间、实测值、测试仪编号和测试员。"),
    ("h1", "5. 异常处理"),
    ("p", "发生击穿、飞弧或测试仪报警时，立即切断高压并挂“不合格”标识。不得重复测试同一产品以试图获得合格结果；应隔离产品并通知质量部。"),
]

BOM_X100_BLOCKS = [
    ("h1", "1. 说明"),
    ("p", "本清单为 X-100 便携式扭矩校验仪的整机物料清单，适用于样机与小批量生产。替代料必须经过质量部与研发部确认。"),
    ("h1", "2. 物料清单"),
    ("table", ["序号", "物料编码", "名称", "规格/型号", "数量"], [
        ["1", "ST-ME-001", "铝合金上壳", "X100-ALK-01", "1"],
        ["2", "ST-ME-002", "铝合金下壳", "X100-ALK-02", "1"],
        ["3", "ST-PCB-010", "主控板", "X100-MB-3.2", "1"],
        ["4", "ST-SEN-003", "扭矩传感器", "TQ-100-0.5", "1"],
        ["5", "ST-DIS-002", "液晶显示模组", "2.4 inch IPS", "1"],
        ["6", "ST-BAT-004", "锂离子电池", "3.7 V 2600 mAh", "1"],
        ["7", "ST-BT-005", "蓝牙模组", "BT5.0-UART", "1"],
        ["8", "ST-ME-010", "按键膜", "X100-KEY-01", "1"],
        ["9", "ST-ACC-006", "USB-C 防水塞", "IP54", "1"],
        ["10", "ST-ACC-007", "携带箱", "EVA-210", "1"],
    ], [1.2, 3.2, 4.0, 4.5, 2.0], "表 1  X-100 物料清单"),
    ("h1", "3. 关键件要求"),
    ("p", "扭矩传感器、主控板和电池为关键件，进货检验按《来料检验规范》执行。传感器应提供灵敏度与温度补偿报告；电池应提供 UN38.3 运输鉴定报告。"),
]

BOM_X200_BLOCKS = [
    ("h1", "1. 说明"),
    ("p", "本清单为 X-200 台式多通道扭矩校验系统的整机物料清单。系统为固定式设备，不含电池与携带箱。"),
    ("h1", "2. 物料清单"),
    ("table", ["序号", "物料编码", "名称", "规格/型号", "数量"], [
        ["1", "ST-X200-FR", "机架", "X200-FR-01", "1"],
        ["2", "ST-X200-TS", "12 英寸触摸屏", "IPS-12-10P", "1"],
        ["3", "ST-X200-MB", "主控板", "X200-MB-1.0", "1"],
        ["4", "ST-X200-CH", "四通道采集卡", "AD4-24BIT", "1"],
        ["5", "ST-X200-SEN", "扭矩传感器", "TQ-2000-0.2", "4"],
        ["6", "ST-X200-PSU", "开关电源", "24 V 12.5 A", "1"],
        ["7", "ST-X200-FAN", "散热风扇", "80 mm", "2"],
        ["8", "ST-X200-CASE", "钣金外壳", "X200-CASE-02", "1"],
    ], [1.2, 3.2, 4.0, 4.5, 2.0], "表 1  X-200 物料清单"),
    ("h1", "3. 装配要求"),
    ("p", "四通道采集卡与传感器必须按通道号一一对应，装配后执行零点与满量程自检。机架接地端子应使用不小于 4 mm² 的黄绿线连接。"),
]

CAL_TEMPLATE_BLOCKS = [
    ("h1", "1. 适用范围"),
    ("p", "本模板适用于 X-100、X-100S 扭矩校验仪的周期校准记录。校准周期为 12 个月，或累计使用 5000 次后提前校准。"),
    ("h1", "2. 校准条件"),
    ("table", ["项目", "要求"], [
        ["环境温度", "20 ± 5 °C"],
        ["相对湿度", "45%–75%"],
        ["标准扭矩源", "不确定度优于 0.1%"],
        ["预热时间", "≥ 30 min"],
        ["校准点", "20%、40%、60%、80%、100% FS"],
    ], [5.0, 8.0], "表 1  校准条件"),
    ("h1", "3. 记录字段"),
    ("p", "记录应包含设备型号、序列号、校准日期、校准员、标准器编号、环境温湿度、各校准点的标准值、示值、误差、重复性和结论。"),
    ("table", ["校准点", "标准值（N·m）", "示值（N·m）", "误差（%）", "结论"], [
        ["20%", "", "", "", ""],
        ["40%", "", "", "", ""],
        ["60%", "", "", "", ""],
        ["80%", "", "", "", ""],
        ["100%", "", "", "", ""],
    ], [3.0, 4.0, 4.0, 3.0, 3.0], "表 2  校准记录"),
    ("h1", "4. 判定"),
    ("p", "各点示值误差不超过 ±0.5% FS、重复性不超过 0.3% 时判定合格。任一点超差应调整或维修后重新校准；维修后必须重新执行全部校准点。"),
]

CAL_OLD_BLOCKS = [
    ("h1", "1. 说明"),
    ("p", "本记录模板为 2023 年旧版，已被 2026 版《校准记录模板》替代。本文件仅用于历史记录追溯，不得用于新设备校准。"),
    ("h1", "2. 旧版要求"),
    ("p", "旧版校准周期为 6 个月，校准点为量程的 10%、30%、50%、70%、90%。旧版判定限为示值误差不超过 ±0.8% FS。"),
    ("table", ["校准点", "标准值", "示值", "误差", "结论"], [
        ["10%", "", "", "", ""],
        ["30%", "", "", "", ""],
        ["50%", "", "", "", ""],
        ["70%", "", "", "", ""],
        ["90%", "", "", "", ""],
    ], [3.0, 4.0, 4.0, 3.0, 3.0], "表 1  旧版校准记录"),
    ("h1", "3. 存档说明"),
    ("p", "旧版记录仅供存档，不得复制用于新批次。新设备一律使用 2026 版模板，校准周期为 12 个月。"),
]

PACKAGING_BLOCKS = [
    ("h1", "1. 适用范围"),
    ("p", "本规范适用于 X-100、X-100S 便携式扭矩校验仪的包装、标识、堆码与运输。X-200 台式系统使用单独的木箱包装规范。"),
    ("h1", "2. 包装材料"),
    ("bullets", [
        "外箱：五层瓦楞纸箱，内尺寸 320 × 220 × 120 mm。",
        "内衬：EVA 成型内衬，厚度不小于 20 mm。",
        "防潮：每箱放置 10 g 干燥剂两包。",
        "标识：防潮、向上、易碎、堆码层数标识。",
    ]),
    ("h1", "3. 堆码与运输"),
    ("p", "仓库堆码不得超过 5 层，单箱毛重不超过 3.5 kg。运输过程中应避免雨淋、倒置和剧烈冲击；装卸时不得抛掷。"),
    ("table", ["项目", "要求"], [
        ["堆码层数", "≤ 5 层"],
        ["跌落高度", "1.0 m，六面各一次"],
        ["振动", "10–55 Hz，30 min"],
        ["环境湿度", "20%–80% RH"],
        ["存储温度", "-20 °C ~ 60 °C"],
    ], [5.0, 8.0], "表 1  包装运输要求"),
    ("h1", "4. 标识与随箱文件"),
    ("p", "外箱标签应包含产品型号、序列号、生产日期、毛重、箱号和生产批次。随箱文件包括说明书、合格证、校准证书和保修卡。"),
]

FIR_X100_BLOCKS = [
    ("h1", "1. 基本信息"),
    ("table", ["项目", "内容"], [
        ["产品型号", "X-100"],
        ["序列号", "X100-2026-0912-001"],
        ["生产批次", "2026-09-A"],
        ["检验依据", "STARLAN-QS-001 V2.1"],
        ["检验日期", "2026-09-12"],
        ["检验员", "QA-017"],
    ], [4.0, 9.0], "表 1  样品信息"),
    ("h1", "2. 检验结果"),
    ("table", ["项目", "实测值", "判定"], [
        ["外观", "无裂纹、无锈蚀、标识清晰", "合格"],
        ["外形尺寸", "209.8 × 95.1 × 42.0 mm", "合格"],
        ["整机重量", "0.681 kg", "合格"],
        ["扭矩示值误差", "+0.18% FS", "合格"],
        ["重复性", "0.12%", "合格"],
        ["绝缘电阻", "620 MΩ（500 VDC）", "合格"],
        ["耐压", "2000 VAC / 60 s，无击穿", "合格"],
        ["防护等级", "IP54", "合格"],
    ], [4.0, 7.0, 3.0], "表 2  检验结果"),
    ("h1", "3. 结论"),
    ("p", "该样品依据 STARLAN-QS-001 V2.1 检验，所检项目全部合格，判定为出厂检验合格。本记录为虚构演示样例，不对应真实产品。"),
]

FIR_X200_BLOCKS = [
    ("h1", "1. 基本信息"),
    ("table", ["项目", "内容"], [
        ["产品型号", "X-200"],
        ["序列号", "X200-2026-0903-003"],
        ["生产批次", "2026-09-B"],
        ["检验依据", "X-200 出厂检验作业指导书 WI-QA-022"],
        ["检验日期", "2026-09-03"],
        ["检验员", "QA-021"],
    ], [4.0, 9.0], "表 1  样品信息"),
    ("h1", "2. 检验结果"),
    ("table", ["项目", "实测值", "判定"], [
        ["外观", "无划伤、屏幕无亮点", "合格"],
        ["四通道零点", "CH1–CH4 均 ≤0.02% FS", "合格"],
        ["扭矩示值误差", "CH1 +0.08% FS；CH2 +0.06% FS", "合格"],
        ["通道一致性", "≤0.05% FS", "合格"],
        ["以太网通信", "Modbus TCP 连接正常", "合格"],
        ["接地电阻", "0.8 Ω", "合格"],
        ["触摸屏", "全区域响应正常", "合格"],
    ], [4.0, 7.5, 3.0], "表 2  检验结果"),
    ("h1", "3. 结论"),
    ("p", "该样品依据 X-200 出厂检验作业指导书检验，所检项目全部合格。本记录为虚构演示样例。"),
]

IQC_BLOCKS = [
    ("h1", "1. 目的"),
    ("p", "本规范规定 X-100 系列产品用 PCBA、铝合金外壳、锂离子电池和充电器的来料检验要求、抽样方案与判定规则。"),
    ("h1", "2. 检验项目"),
    ("table", ["物料", "检验项目", "要求", "抽样"], [
        ["PCBA", "外观、通电、固件版本", "无虚焊；版本符合 BOM", "AQL 1.0"],
        ["铝合金外壳", "尺寸、表面、螺纹", "无裂纹；尺寸符合图纸", "AQL 1.0"],
        ["锂离子电池", "电压、内阻、外观", "3.7 V；内阻 ≤80 mΩ", "AQL 0.65"],
        ["充电器", "输出电压、标识", "5 V ±0.25 V；标识完整", "AQL 1.0"],
    ], [3.2, 5.0, 6.0, 2.5], "表 1  来料检验项目"),
    ("h1", "3. 判定与记录"),
    ("p", "关键项不合格时整批拒收；一般项按 AQL 判定。检验记录保存 3 年，包含供应商、批次、抽样数、不合格数和处理结论。"),
]

CLEANING_BLOCKS = [
    ("h1", "1. 适用范围"),
    ("p", "本规程适用于 M-3、M-3 Pro 装配台及周边工装的日常清洁。清洁前必须停机、切断伺服使能并释放气源。"),
    ("h1", "2. 清洁工具与材料"),
    ("bullets", [
        "无绒布、软毛刷、吸尘器。",
        "中性清洁剂（屏幕与外壳）。",
        "无水乙醇（导轨与金属夹具）。",
        "禁止使用 WD-40、香蕉水、丙酮和高压水枪。",
    ]),
    ("h1", "3. 清洁步骤"),
    ("p", "每班结束后清理台面与夹具上的金属屑，用吸尘器清理导轨缝隙，用无绒布擦拭触摸屏。导轨清洁只能使用无水乙醇，清洁后按保养表补充润滑脂。"),
    ("p", "清洁过程中不得移动安全光栅、急停按钮和限位开关。清洁完成后执行空运行，确认无异物干涉。"),
    ("h1", "4. 记录"),
    ("p", "清洁记录填入《设备日常点检表》，包含日期、班次、清洁人、异常情况和处理结果。"),
]

ESD_BLOCKS = [
    ("h1", "1. 适用范围"),
    ("p", "本规范适用于 X-100、X-100S、X-200 的 PCBA 装配、调试、返修和包装区域。"),
    ("h1", "2. 防静电要求"),
    ("table", ["项目", "要求"], [
        ["防静电腕带", "1 MΩ 限流电阻，每周测试一次"],
        ["防静电地板", "表面电阻 1×10^6 – 1×10^9 Ω"],
        ["工作台", "台垫接地，接地电阻 <1 Ω"],
        ["环境湿度", "40%–60% RH"],
        ["离子风机", "每日点检，覆盖关键工位"],
    ], [4.0, 9.0], "表 1  防静电要求"),
    ("h1", "3. 操作要求"),
    ("p", "接触 PCBA 前必须佩戴腕带并确认测试合格。PCBA 不得直接放在普通桌面，应使用防静电周转盒。维修工位应配备离子风机和接地台垫。"),
    ("h1", "4. 审核"),
    ("p", "ESD 管理员每月抽查腕带、台垫和地板电阻，记录不符合项并跟踪整改。"),
]

OLD_M3_BLOCKS = [
    ("h1", "1. 说明"),
    ("p", "本操作手册为 M-3 早期版本（V0.8），已被《M-3 操作手册 V1.1》替代。本文件仅用于历史存档，不得用于当前生产。"),
    ("h1", "2. 旧版参数"),
    ("p", "旧版默认锁付扭矩为 6.0 N·m，转速 240 rpm，单件节拍 ≤15 s。旧版未配置安全光栅，仅使用安全门。"),
    ("h1", "3. 旧版操作步骤"),
    ("p", "开机后回原点，选择程序，手动放入工件，双手按下启动按钮。旧版没有首件确认流程，操作员自检后直接批量生产。"),
    ("h1", "4. 存档提示"),
    ("p", "当前生产必须使用 V1.1 手册，默认扭矩 8.0 N·m，并执行安全光栅、急停和首件确认要求。本旧版内容不得作为操作依据。"),
]

SQA_BLOCKS = [
    ("h1", "1. 目的"),
    ("p", "本协议规定星澜精密仪器（虚构）有限公司对关键物料供应商的质量要求，包括样品承认、批量交付、变更通知、异常处理和审核。"),
    ("h1", "2. 质量要求"),
    ("bullets", [
        "关键物料批次合格率不低于 99.5%。",
        "重大变更（材料、工艺、产地、软件版本）应提前 30 天书面通知。",
        "来料不合格时供应商应在 24 h 内响应，5 个工作日内提交 8D 报告。",
        "关键物料应保留批次追溯信息不少于 5 年。",
    ]),
    ("h1", "3. 审核与绩效"),
    ("p", "质量部每年对关键供应商进行一次现场审核；出现重大质量事故时增加专项审核。供应商绩效按批次合格率、响应速度、8D 关闭率每季度评分。"),
    ("h1", "4. 变更与召回"),
    ("p", "未经书面批准的变更，星澜有权拒收相关批次。已交付产品发生质量问题时，供应商应配合追溯、隔离和召回，并承担由此产生的直接损失。"),
]

# ---------------------------------------------------------------------------
# Document specs
# ---------------------------------------------------------------------------

DOCS: list[dict[str, Any]] = [
    {
        "id": "A01", "file": "STARLAN-X100-product-manual-v1.2.docx", "format": "docx",
        "folder": "product", "category": "产品说明书", "title": "X-100 便携式扭矩校验仪产品说明书",
        "code": "STARLAN-PM-X100-001", "version": "V1.2", "date": "2025-11-20",
        "status": "受控", "scope": "X-100 便携式扭矩校验仪",
        "summary": "X-100 便携式扭矩校验仪，额定功率 5 W，IP54，Bluetooth 5.0。",
        "blocks": X100_COMMON,
    },
    {
        "id": "A02", "file": "STARLAN-X100S-product-manual-v1.2.docx", "format": "docx",
        "folder": "product", "category": "产品说明书", "title": "X-100S 便携式扭矩校验仪产品说明书",
        "code": "STARLAN-PM-X100S-001", "version": "V1.2", "date": "2026-02-18",
        "status": "受控", "scope": "X-100S 便携式扭矩校验仪",
        "summary": "X-100S 便携式扭矩校验仪，额定功率 8 W，IP67，Wi-Fi 6 + Bluetooth 5.2。",
        "blocks": X100S_BLOCKS,
    },
    {
        "id": "A03", "file": "STARLAN-X200-product-manual-v1.0.pdf", "format": "pdf",
        "folder": "product", "category": "产品说明书", "title": "X-200 台式多通道扭矩校验系统产品说明书",
        "code": "STARLAN-PM-X200-001", "version": "V1.0", "date": "2026-05-09",
        "status": "受控", "scope": "X-200 台式多通道扭矩校验系统",
        "summary": "X-200 台式四通道系统，5–2000 N·m，±0.2% FS，220 V 供电，IP20。",
        "blocks": X200_BLOCKS,
    },
    {
        "id": "B04", "file": "STARLAN-QS-001-inspection-spec-v2.0.pdf", "format": "pdf",
        "folder": "quality", "category": "质量检验规范", "title": "X-100 系列成品检验规范 V2.0",
        "code": "STARLAN-QS-001", "version": "V2.0", "date": "2025-03-10",
        "status": "历史版本（被 V2.1 替代）", "scope": "X-100 / X-100S 成品",
        "summary": "V2.0 旧版规范：耐压 1500 VAC，绝缘电阻 ≥100 MΩ，误差 ±0.5% FS。",
        "blocks": QS20_BLOCKS,
    },
    {
        "id": "B05", "file": "STARLAN-QS-001-inspection-spec-v2.1.docx", "format": "docx",
        "folder": "quality", "category": "质量检验规范", "title": "X-100 系列成品检验规范 V2.1",
        "code": "STARLAN-QS-001", "version": "V2.1", "date": "2026-08-15",
        "status": "受控（现行）", "scope": "X-100 / X-100S 成品",
        "summary": "V2.1 现行规范：耐压 2000 VAC，绝缘电阻 ≥500 MΩ，误差 ±0.4% FS。",
        "blocks": QS21_BLOCKS,
    },
    {
        "id": "C06", "file": "STARLAN-M3-operation-manual-v1.1.pdf", "format": "pdf",
        "folder": "equipment", "category": "设备操作手册", "title": "M-3 半自动扭矩装配台操作手册",
        "code": "STARLAN-EQ-M3-001", "version": "V1.1", "date": "2026-01-12",
        "status": "受控", "scope": "M-3 半自动扭矩装配台",
        "summary": "M-3 默认扭矩 8.0 N·m，转速 300 rpm，节拍 ≤12 s。",
        "blocks": M3_BLOCKS,
    },
    {
        "id": "C07", "file": "STARLAN-M3-Pro-operation-manual-v1.0.docx", "format": "docx",
        "folder": "equipment", "category": "设备操作手册", "title": "M-3 Pro 全自动扭矩装配台操作手册",
        "code": "STARLAN-EQ-M3P-001", "version": "V1.0", "date": "2026-06-30",
        "status": "受控", "scope": "M-3 Pro 全自动扭矩装配台",
        "summary": "M-3 Pro 默认扭矩 12.0 N·m，转速 450 rpm，节拍 ≤8 s；提及 X-900 扩展型号。",
        "blocks": M3PRO_BLOCKS,
    },
    {
        "id": "C08", "file": "STARLAN-equipment-maintenance-procedure-v1.3.docx", "format": "docx",
        "folder": "equipment", "category": "设备维护规程", "title": "M-3 / M-3 Pro 扭矩装配台维护规程",
        "code": "STARLAN-EQ-MNT-001", "version": "V1.3", "date": "2026-07-08",
        "status": "受控", "scope": "M-3 / M-3 Pro",
        "summary": "日常/周/月保养与 6 个月扭矩校准要求。",
        "blocks": MAINT_BLOCKS,
    },
    {
        "id": "D09", "file": "STARLAN-SAFE-001-assembly-safety-procedure-v1.1.pdf", "format": "pdf",
        "folder": "safety", "category": "安全操作规程", "title": "装配工位安全操作规程",
        "code": "STARLAN-SAFE-001", "version": "V1.1", "date": "2026-03-05",
        "status": "受控", "scope": "M-3 / M-3 Pro 装配工位",
        "summary": "装配工位 PPE、LOTO、急停后等待 3 s、换型安全要求。",
        "blocks": SAFE_ASSEMBLY_BLOCKS,
    },
    {
        "id": "D10", "file": "STARLAN-SAFE-002-electrical-test-safety-v1.0.docx", "format": "docx",
        "folder": "safety", "category": "安全操作规程", "title": "电气测试安全操作规程",
        "code": "STARLAN-SAFE-002", "version": "V1.0", "date": "2026-03-20",
        "status": "受控", "scope": "耐压 / 绝缘 / 接地测试",
        "summary": "高压测试双人监护、耐压后放电不少于 5 分钟。",
        "blocks": SAFE_ELECTRICAL_BLOCKS,
    },
    {
        "id": "D11", "file": "STARLAN-BOM-X100-v1.2.docx", "format": "docx",
        "folder": "bom", "category": "物料清单", "title": "X-100 整机物料清单",
        "code": "STARLAN-BOM-X100-001", "version": "V1.2", "date": "2026-04-11",
        "status": "受控", "scope": "X-100 整机",
        "summary": "X-100 十项关键物料与关键件要求。",
        "blocks": BOM_X100_BLOCKS,
    },
    {
        "id": "D12", "file": "STARLAN-BOM-X200-v1.0.pdf", "format": "pdf",
        "folder": "bom", "category": "物料清单", "title": "X-200 整机物料清单",
        "code": "STARLAN-BOM-X200-001", "version": "V1.0", "date": "2026-05-15",
        "status": "受控", "scope": "X-200 整机",
        "summary": "X-200 八项关键物料与装配接地要求。",
        "blocks": BOM_X200_BLOCKS,
    },
    {
        "id": "D13", "file": "STARLAN-CAL-record-template-v1.0.docx", "format": "docx",
        "folder": "calibration", "category": "校准记录模板", "title": "X-100 系列扭矩校准记录模板",
        "code": "STARLAN-CAL-FORM-001", "version": "V1.0", "date": "2026-01-05",
        "status": "受控", "scope": "X-100 / X-100S",
        "summary": "校准周期 12 个月，五点校准，误差限 ±0.5% FS。",
        "blocks": CAL_TEMPLATE_BLOCKS,
    },
    {
        "id": "D14", "file": "STARLAN-CAL-record-template-v0.9-obsolete.pdf", "format": "pdf",
        "folder": "calibration", "category": "校准记录模板", "title": "X-100 系列扭矩校准记录模板（旧版）",
        "code": "STARLAN-CAL-FORM-000", "version": "V0.9", "date": "2023-06-01",
        "status": "已废止", "scope": "历史记录存档",
        "summary": "已废止旧模板：校准周期 6 个月，误差限 ±0.8% FS。",
        "blocks": CAL_OLD_BLOCKS, "obsolete": True,
    },
    {
        "id": "D15", "file": "STARLAN-PKG-001-packaging-transport-spec-v1.1.pdf", "format": "pdf",
        "folder": "packaging", "category": "包装运输规范", "title": "X-100 系列包装与运输规范",
        "code": "STARLAN-PKG-001", "version": "V1.1", "date": "2026-02-02",
        "status": "受控", "scope": "X-100 / X-100S",
        "summary": "堆码不超过 5 层，1.0 m 六面跌落，防潮防倒置。",
        "blocks": PACKAGING_BLOCKS,
    },
    {
        "id": "D16", "file": "STARLAN-FIR-X100-sample-v1.0.docx", "format": "docx",
        "folder": "reports", "category": "出厂检验报告样例", "title": "X-100 出厂检验报告样例",
        "code": "STARLAN-FIR-X100-001", "version": "V1.0", "date": "2026-09-12",
        "status": "样例（虚构）", "scope": "X-100 样品 X100-2026-0912-001",
        "summary": "X-100 样品检验记录：绝缘 620 MΩ，耐压 2000 VAC，结论合格。",
        "blocks": FIR_X100_BLOCKS,
    },
    {
        "id": "D17", "file": "STARLAN-FIR-X200-sample-v1.0.pdf", "format": "pdf",
        "folder": "reports", "category": "出厂检验报告样例", "title": "X-200 出厂检验报告样例",
        "code": "STARLAN-FIR-X200-001", "version": "V1.0", "date": "2026-09-03",
        "status": "样例（虚构）", "scope": "X-200 样品 X200-2026-0903-003",
        "summary": "X-200 四通道零点、扭矩误差、以太网通信检验记录。",
        "blocks": FIR_X200_BLOCKS,
    },
    {
        "id": "D18", "file": "STARLAN-IQC-001-incoming-material-spec-v1.2.docx", "format": "docx",
        "folder": "quality", "category": "来料检验规范", "title": "关键物料来料检验规范",
        "code": "STARLAN-IQC-001", "version": "V1.2", "date": "2026-04-22",
        "status": "受控", "scope": "PCBA / 外壳 / 电池 / 充电器",
        "summary": "关键物料抽样 AQL 0.65/1.0，电池内阻 ≤80 mΩ。",
        "blocks": IQC_BLOCKS,
    },
    {
        "id": "D19", "file": "STARLAN-equipment-cleaning-sop-v1.0.pdf", "format": "pdf",
        "folder": "equipment", "category": "设备清洁规程", "title": "M-3 系列装配台清洁规程",
        "code": "STARLAN-EQ-CLN-001", "version": "V1.0", "date": "2026-03-18",
        "status": "受控", "scope": "M-3 / M-3 Pro",
        "summary": "导轨只用无水乙醇，禁止 WD-40/高压水枪。",
        "blocks": CLEANING_BLOCKS,
    },
    {
        "id": "D20", "file": "STARLAN-ESD-control-spec-v1.0.docx", "format": "docx",
        "folder": "safety", "category": "ESD 控制规范", "title": "电子装配防静电控制规范",
        "code": "STARLAN-ESD-001", "version": "V1.0", "date": "2026-05-06",
        "status": "受控", "scope": "PCBA 装配 / 调试 / 返修",
        "summary": "腕带 1 MΩ，每周测试；地板 1e6–1e9 Ω。",
        "blocks": ESD_BLOCKS,
    },
    {
        "id": "D21", "file": "STARLAN-M3-old-operation-procedure-v0.8-obsolete.pdf", "format": "pdf",
        "folder": "archive", "category": "历史操作规程", "title": "M-3 半自动扭矩装配台操作规程（旧版）",
        "code": "STARLAN-EQ-M3-000", "version": "V0.8", "date": "2023-08-15",
        "status": "已废止", "scope": "历史存档",
        "summary": "已废止旧规程：默认扭矩 6.0 N·m，无安全光栅与首件确认。",
        "blocks": OLD_M3_BLOCKS, "obsolete": True,
    },
    {
        "id": "D22", "file": "STARLAN-SQA-001-supplier-quality-agreement-v1.1.docx", "format": "docx",
        "folder": "supplier", "category": "供应商质量协议", "title": "关键物料供应商质量协议",
        "code": "STARLAN-SQA-001", "version": "V1.1", "date": "2026-06-12",
        "status": "受控", "scope": "关键物料供应商",
        "summary": "批次合格率 99.5%，重大变更提前 30 天通知，24 h 响应。",
        "blocks": SQA_BLOCKS,
    },
]


# ---------------------------------------------------------------------------
# Build
# ---------------------------------------------------------------------------


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def convert_pdf(docx_path: Path, pdf_path: Path) -> None:
    if not NODE.exists() or not LIBREOFFICE_CLI.exists():
        raise FileNotFoundError("LibreOffice Kit node/cli not found")
    result = subprocess.run(
        [str(NODE), str(LIBREOFFICE_CLI), "convert", "--input", str(docx_path), "--output", str(pdf_path)],
        capture_output=True,
        text=True,
        check=False,
    )
    if result.returncode != 0 or not pdf_path.exists():
        raise RuntimeError(
            f"PDF conversion failed for {docx_path.name}: "
            f"exit={result.returncode} stdout={result.stdout[-500:]} stderr={result.stderr[-500:]}"
        )
    if pdf_path.read_bytes()[:4] != b"%PDF":
        raise RuntimeError(f"not a PDF: {pdf_path}")


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--only", default="", help="comma separated document ids")
    parser.add_argument("--work-dir", type=Path, default=DEFAULT_WORK_DIR)
    args = parser.parse_args()

    selected = {x.strip().upper() for x in args.only.split(",") if x.strip()}
    docs = [d for d in DOCS if not selected or d["id"].upper() in selected]
    if not docs:
        raise SystemExit("no documents selected")

    work_docx = args.work_dir / "docx"
    work_pdf = args.work_dir / "pdf"
    work_docx.mkdir(parents=True, exist_ok=True)
    work_pdf.mkdir(parents=True, exist_ok=True)

    manifest: list[dict[str, Any]] = []
    for spec in docs:
        final_path = DATA_DIR / spec["folder"] / spec["file"]
        final_path.parent.mkdir(parents=True, exist_ok=True)
        docx_path = work_docx / (Path(spec["file"]).stem + ".docx")

        doc = new_document(spec)
        doc.save(docx_path)

        if spec["format"] == "pdf":
            pdf_path = work_pdf / (Path(spec["file"]).stem + ".pdf")
            convert_pdf(docx_path, pdf_path)
            shutil.copy2(pdf_path, final_path)
        else:
            shutil.copy2(docx_path, final_path)

        record = {
            "id": spec["id"],
            "path": final_path.relative_to(REPO_ROOT).as_posix(),
            "format": spec["format"],
            "category": spec["category"],
            "title": spec["title"],
            "code": spec["code"],
            "version": spec["version"],
            "date": spec["date"],
            "status": spec["status"],
            "scope": spec["scope"],
            "summary": spec["summary"],
            "obsolete": bool(spec.get("obsolete", False)),
            "bytes": final_path.stat().st_size,
            "sha256": sha256(final_path),
        }
        manifest.append(record)
        print(f"[OK] {spec['id']} {spec['format']:4} {record['path']} ({record['bytes']} bytes)")

    manifest_path = DATA_DIR / "manifest.json"
    if selected and manifest_path.exists():
        existing = json.loads(manifest_path.read_text(encoding="utf-8"))
        by_id = {item["id"]: item for item in existing}
        by_id.update({item["id"]: item for item in manifest})
        manifest = [by_id[key] for key in sorted(by_id)]
    manifest_path.write_text(json.dumps(manifest, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(f"[OK] manifest: {manifest_path} ({len(manifest)} documents)")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
